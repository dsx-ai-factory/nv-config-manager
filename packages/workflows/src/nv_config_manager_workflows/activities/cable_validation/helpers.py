# SPDX-FileCopyrightText: Copyright (c) 2024-2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Pure validation and formatting helpers for cable-validation activities."""

from __future__ import annotations

import base64
import csv
import io
from collections.abc import Iterable
from typing import Any

import netaddr
import pandas as pd
from nv_config_manager_dcim import CableStatus
from nv_config_manager_dcim.workflow_models import NetworkDeviceData
from openpyxl.utils import get_column_letter
from py_markdown_table.markdown_table import markdown_table
from temporalio.exceptions import ApplicationError

from nv_config_manager_workflows.activities.cable_validation.models import (
    DETAIL_SHEET_NAME,
    EXCEL_MIME_TYPE,
    HOST_SUMMARY_COLUMNS,
    HOST_SUMMARY_SHEET_NAME,
    INCORRECT_CABLING_PREFIX,
    LINK_DOWN_MSG,
    LINK_UP_NO_NEIGHBOR_MSG,
    UNEXPECTED_CONNECTION_MSG,
    CableValidationResultData,
    CableValidationRow,
    InvalidCable,
    _escape_formula,
)
from nv_config_manager_workflows.clients.device.models import (
    DeviceArpTable,
    DeviceMacTable,
    DeviceNeighborData,
    InterfaceNeighborData,
    format_mac,
    is_mac_address,
)


def _rack_position_str(rack: str | None, position: int | None) -> str | None:
    """Return rack position string (e.g. 'rack1:u42') or None if rack/position missing."""
    if rack and position is not None:
        return f"{rack}:u{position}"
    return None


def _display_or_macs(primary: str | None, macs: list[str] | None) -> str | None:
    """Return primary if set, else comma-joined macs or None (for Actual End Device/Port)."""
    return primary or (",".join(macs) if macs else None)


def _as_mac_address(s: str) -> str | None:
    try:
        return str(netaddr.EUI(s))
    except netaddr.core.AddrFormatError:
        return None


def _mac_matches_with_offset(actual_mac: str, expected_mac: str) -> bool:
    """
    Check if actual MAC matches expected MAC or expected MAC + 0x10 offset.

    DPUs sometimes advertise their MAC over LLDP with an offset of 0x10 when in NIC mode.
    This function checks both the original expected MAC and the offset version.

    Args:
        actual_mac: The MAC address received via LLDP
        expected_mac: The intended/expected MAC address

    Returns:
        True if actual_mac matches expected_mac or expected_mac + 0x10
    """
    try:
        actual_eui = netaddr.EUI(actual_mac)
        expected_eui = netaddr.EUI(expected_mac)

        # Check exact match first
        if actual_eui == expected_eui:
            return True

        # Check if actual MAC matches expected MAC + 0x10 offset
        offset_eui = netaddr.EUI(expected_eui.value + 0x10)
        return bool(actual_eui == offset_eui)

    except (netaddr.core.AddrFormatError, ValueError):
        return False


def _validate_neighbor_has_required_fields(
    intended_neighbor: InterfaceNeighborData, device: NetworkDeviceData
) -> None:
    """Validate that intended neighbor has required role and name fields."""
    if not intended_neighbor.device_role:
        raise ApplicationError(
            f"No role information for intended neighbor {intended_neighbor} on device {device}"
        )
    if not intended_neighbor.device_name:
        raise ApplicationError(
            f"No name information for intended neighbor {intended_neighbor} on device {device}"
        )


def _check_link_state(
    intended_interface: str,
    actual: DeviceNeighborData,
    intended_neighbor: InterfaceNeighborData,
) -> InvalidCable | None:
    """Check link state and return InvalidCable if link is down."""
    if not actual.link_states.get(intended_interface, False):
        return InvalidCable(
            intended=intended_neighbor,
            actual=InterfaceNeighborData(
                link_up=False,
                ts_info=actual.ts_info.get(intended_interface),
            ),
        )
    return None


def _check_interface_mac_match(
    actual_neighbor_interface: str | None, expected_mac: str | None
) -> bool:
    """Check if actual neighbor interface MAC matches expected MAC."""
    if not actual_neighbor_interface or not expected_mac:
        return False

    neighbor_if_mac = _as_mac_address(actual_neighbor_interface)
    if neighbor_if_mac:
        return _mac_matches_with_offset(neighbor_if_mac, expected_mac)
    return False


def _check_device_mac_match(actual_neighbor_device: str | None, expected_mac: str | None) -> bool:
    """Check if actual neighbor device MAC matches expected MAC."""
    if not actual_neighbor_device or not expected_mac:
        return False

    neighbor_device_mac = _as_mac_address(actual_neighbor_device)
    if neighbor_device_mac:
        return _mac_matches_with_offset(neighbor_device_mac, expected_mac)
    return False


def _check_standard_lldp_match(
    intended_neighbor_interface: str | None,
    intended_neighbor_device: str | None,
    actual_neighbor_interface: str | None,
    actual_neighbor_device: str | None,
) -> bool:
    """Check if LLDP data matches using standard interface and device names."""
    return (
        intended_neighbor_interface == actual_neighbor_interface
        and intended_neighbor_device == actual_neighbor_device
    )


def _check_lldp_validity(
    intended_interface: str,
    intended_neighbor: InterfaceNeighborData,
    actual: DeviceNeighborData,
) -> bool:
    """Check if LLDP data is valid."""
    if intended_interface not in actual.neighbors:
        return False

    actual_neighbor = actual.neighbors[intended_interface]

    # Normalize strings to lowercase for comparison
    intended_neighbor_interface = intended_neighbor.name.lower() if intended_neighbor.name else None
    intended_neighbor_device = (
        intended_neighbor.device_name.lower() if intended_neighbor.device_name else None
    )
    actual_neighbor_interface = actual_neighbor.name.lower() if actual_neighbor.name else None
    actual_neighbor_device = (
        actual_neighbor.device_name.lower() if actual_neighbor.device_name else None
    )

    expected_mac = intended_neighbor.macs[0] if intended_neighbor.macs else None

    # Device sent the interface mac instead of name
    if _check_interface_mac_match(actual_neighbor_interface, expected_mac):
        return True

    # Device sent mac address instead of device name
    if _check_device_mac_match(actual_neighbor_device, expected_mac):
        return True

    # Standard LLDP check
    return _check_standard_lldp_match(
        intended_neighbor_interface,
        intended_neighbor_device,
        actual_neighbor_interface,
        actual_neighbor_device,
    )


def _check_mac_validity(
    intended_interface: str,
    intended_neighbor: InterfaceNeighborData,
    mac_table: DeviceMacTable,
    arp_table: DeviceArpTable,
) -> bool:
    """Check if MAC address is valid."""
    if not intended_neighbor.macs:
        return False

    mac = intended_neighbor.macs[0]
    return (mac in mac_table.by_mac and mac_table.by_mac[mac].interface == intended_interface) or (
        mac in arp_table.interface_to_mac.get(intended_interface, [])
    )


def _build_actual_neighbor_data(
    intended_interface: str,
    intended_neighbor: InterfaceNeighborData,
    actual: DeviceNeighborData,
    mac_table: DeviceMacTable,
    arp_table: DeviceArpTable,
) -> InterfaceNeighborData:
    """Build actual neighbor data for error reporting."""
    mac_neighbor = InterfaceNeighborData(
        macs=(
            mac_table.by_interface[intended_interface]
            if intended_interface in mac_table.by_interface
            and mac_table.by_interface[intended_interface]
            else arp_table.interface_to_mac.get(intended_interface, [])
        ),
        link_up=actual.link_states.get(intended_interface, False),
    )

    # Prefer MAC error message if MAC data present
    if intended_neighbor.macs and (
        mac_table.by_interface.get(intended_interface)
        or arp_table.interface_to_mac.get(intended_interface)
    ):
        return mac_neighbor

    if intended_interface in actual.neighbors:
        return actual.neighbors[intended_interface]

    return mac_neighbor


def _process_intended_interface(
    intended_interface: str,
    intended_neighbor: InterfaceNeighborData,
    intended: DeviceNeighborData,
    actual: DeviceNeighborData,
    device: NetworkDeviceData,
    mac_table: DeviceMacTable,
    arp_table: DeviceArpTable,
    ignore_no_neighbor: bool,
) -> tuple[InvalidCable | None, CableStatus | None]:
    """Process one intended interface and return its issue and cable status."""
    if intended_interface in intended.ignore:
        return None, None

    _validate_neighbor_has_required_fields(intended_neighbor, device)

    # Check link state
    link_state_result = _check_link_state(intended_interface, actual, intended_neighbor)
    if link_state_result:
        return link_state_result, CableStatus.DISCONNECTED

    # Skip further checks if link state only validation is requested
    if intended_interface in intended.link_state_only:
        return None, CableStatus.CONNECTED

    # Check LLDP and MAC validity
    lldp_valid = _check_lldp_validity(intended_interface, intended_neighbor, actual)
    mac_valid = _check_mac_validity(intended_interface, intended_neighbor, mac_table, arp_table)

    # If either is valid, the cable is valid
    if mac_valid or lldp_valid:
        return None, CableStatus.CONNECTED

    # Build error data
    actual_neighbor = _build_actual_neighbor_data(
        intended_interface, intended_neighbor, actual, mac_table, arp_table
    )

    status: CableStatus | None = CableStatus.INVALID
    if ignore_no_neighbor and not actual_neighbor.name and not actual_neighbor.macs:
        status = None

    return (
        InvalidCable(
            intended=intended_neighbor,
            actual=actual_neighbor,
        ),
        status,
    )


def _find_unexpected_neighbors(
    intended: DeviceNeighborData,
    actual: DeviceNeighborData,
) -> dict[str, InvalidCable]:
    """Find neighbors in actual that aren't in intended."""
    unexpected = {}

    for actual_interface, actual_neighbor in actual.neighbors.items():
        if actual_interface in intended.ignore:
            continue

        if (
            (actual_neighbor.device_name or actual_neighbor.macs or actual_neighbor.device_serial)
            and actual_interface not in intended.neighbors
            and actual_interface
        ):
            unexpected[actual_interface] = InvalidCable(actual=actual_neighbor)

    return unexpected


def _row_unexpected_connection(
    device_name: str,
    interface_name: str,
    device_rack_pos: str | None,
    actual: InterfaceNeighborData,
    actual_device_name: str | None,
) -> CableValidationRow:
    """Build row for unexpected connection (actual neighbor not in intended)."""
    return CableValidationRow(  # type: ignore[call-arg]
        start_device=device_name,
        start_port=interface_name,
        start_rack=device_rack_pos,
        intended_end_device=None,
        intended_end_port=None,
        intended_end_rack=None,
        actual_end_device=_display_or_macs(actual_device_name, actual.macs),
        actual_end_port=_display_or_macs(actual.name, actual.macs),
        issue="Unexpected connection found",
        troubleshooting_info=actual.ts_info,
    )


def _row_link_down(
    device_name: str,
    interface_name: str,
    device_rack_pos: str | None,
    intended_device_name: str | None,
    intended: InterfaceNeighborData,
    actual: InterfaceNeighborData,
) -> CableValidationRow:
    """Build row for link down."""
    return CableValidationRow(  # type: ignore[call-arg]
        start_device=device_name,
        start_port=interface_name,
        start_rack=device_rack_pos,
        intended_end_device=intended_device_name,
        intended_end_port=intended.name,
        intended_end_rack=_rack_position_str(intended.device_rack, intended.device_position),
        actual_end_device=None,
        actual_end_port=None,
        issue="Link is down.",
        troubleshooting_info=actual.ts_info,
    )


def _row_mac_mismatch(
    device_name: str,
    interface_name: str,
    device_rack_pos: str | None,
    intended_device_name: str | None,
    intended: InterfaceNeighborData,
    actual: InterfaceNeighborData,
    actual_device_name: str | None,
) -> CableValidationRow:
    """Build row for incorrect cabling based on MAC comparison."""
    return CableValidationRow(  # type: ignore[call-arg]
        start_device=device_name,
        start_port=interface_name,
        start_rack=device_rack_pos,
        intended_end_device=intended_device_name,
        intended_end_port=intended.name,
        intended_end_rack=_rack_position_str(intended.device_rack, intended.device_position),
        actual_end_device=_display_or_macs(actual_device_name, actual.macs),
        actual_end_port=_display_or_macs(actual.name, actual.macs),
        issue=(
            "Incorrect cabling, actual should match intended. "
            f"Based on expected MAC {intended.macs[0]}*"
        ),
        troubleshooting_info=actual.ts_info,
    )


def _row_lldp_mismatch(
    device_name: str,
    interface_name: str,
    device_rack_pos: str | None,
    intended_device_name: str | None,
    intended: InterfaceNeighborData,
    actual: InterfaceNeighborData,
    actual_device_name: str | None,
) -> CableValidationRow:
    """Build row for incorrect cabling based on LLDP (actual.name set)."""
    end_port = (
        f"{intended.name} ({format_mac(intended.macs[0])})"
        if is_mac_address(actual.name) and intended.macs
        else intended.name
    )
    return CableValidationRow(  # type: ignore[call-arg]
        start_device=device_name,
        start_port=interface_name,
        start_rack=device_rack_pos,
        intended_end_device=intended_device_name,
        intended_end_port=end_port,
        intended_end_rack=_rack_position_str(intended.device_rack, intended.device_position),
        actual_end_device=_display_or_macs(actual_device_name, actual.macs),
        actual_end_port=_display_or_macs(actual.name, actual.macs),
        issue=("Incorrect cabling, actual should match intended. Based on LLDP data"),
        troubleshooting_info=actual.ts_info,
    )


def _row_link_up_no_neighbor(
    device_name: str,
    interface_name: str,
    device_rack_pos: str | None,
    intended_device_name: str | None,
    intended: InterfaceNeighborData,
    actual: InterfaceNeighborData,
) -> CableValidationRow:
    """Build row for link up but no neighbor found."""
    return CableValidationRow(  # type: ignore[call-arg]
        start_device=device_name,
        start_port=interface_name,
        start_rack=device_rack_pos,
        intended_end_device=intended_device_name,
        intended_end_port=intended.name,
        intended_end_rack=_rack_position_str(intended.device_rack, intended.device_position),
        actual_end_device=_display_or_macs(None, actual.macs),
        actual_end_port=_display_or_macs(None, actual.macs),
        issue=LINK_UP_NO_NEIGHBOR_MSG,
        troubleshooting_info=actual.ts_info,
    )


def _format_single_interface_row(
    device_name: str,
    device_rack_pos: str | None,
    interface_name: str,
    interface_result: InvalidCable,
) -> CableValidationRow | None:
    """Convert a single interface validation result into a table row.

    Args:
        device_name: Lowercase device name
        device_rack_pos: Device rack position string (e.g., "rack1:u42")
        interface_name: Interface name
        interface_result: InvalidCable result for this interface

    Returns:
        CableValidationRow with ID set, or None if this should be skipped
    """
    intended = interface_result.intended
    actual = interface_result.actual
    if not actual:
        return None

    intended_device_name = (
        intended.device_name.lower() if intended and intended.device_name else None
    )
    actual_device_name = actual.device_name.lower() if actual.device_name else None

    if not intended:
        row = _row_unexpected_connection(
            device_name, interface_name, device_rack_pos, actual, actual_device_name
        )
    elif actual.link_up is False:
        row = _row_link_down(
            device_name, interface_name, device_rack_pos, intended_device_name, intended, actual
        )
    elif intended.macs and actual.macs:
        row = _row_mac_mismatch(
            device_name,
            interface_name,
            device_rack_pos,
            intended_device_name,
            intended,
            actual,
            actual_device_name,
        )
    elif actual.name:
        row = _row_lldp_mismatch(
            device_name,
            interface_name,
            device_rack_pos,
            intended_device_name,
            intended,
            actual,
            actual_device_name,
        )
    else:
        row = _row_link_up_no_neighbor(
            device_name, interface_name, device_rack_pos, intended_device_name, intended, actual
        )

    row.id_ = CableValidationRow.compute_id(row)
    return row


def _should_skip_link_down_dedup(
    row: CableValidationRow, dedup: set[tuple[str, str]] | None
) -> bool:
    """Return True if row should be skipped (duplicate link-down). Updates dedup."""
    if dedup is None:
        return False
    if row.issue != "Link is down.":
        return False
    intended_device = row.intended_end_device
    intended_port = row.intended_end_port
    if not intended_device or not intended_port:
        return False
    intended_end = (intended_device, intended_port)
    if intended_end in dedup:
        return True
    dedup.add(intended_end)
    return False


def _should_skip_no_neighbor(row: CableValidationRow, ignore_no_neighbor: bool) -> bool:
    """Return True if row should be skipped (link up, no neighbor)."""
    return bool(ignore_no_neighbor and row.issue == LINK_UP_NO_NEIGHBOR_MSG)


def _format_device_result_row(
    device: NetworkDeviceData | None,
    interfaces: dict[str, InvalidCable],
    dedup: set[tuple[str, str]] | None = None,
    ignore_no_neighbor: bool = False,
) -> list[CableValidationRow]:
    """Convert device interface validation results into table rows.

    Args:
        device: NetworkDeviceData object (can be None)
        interfaces: Dictionary of interface names to InvalidCable results
        dedup: Optional set for deduplicating "Link is down" cases
        ignore_no_neighbor: If True, skip "Link is up but no neighbor found" cases

    Returns:
        List of CableValidationRow with IDs set
    """
    device_name = device.name.lower() if device and device.name else "unknown"
    device_rack_pos = _rack_position_str(device.rack, device.position) if device else None

    results: list[CableValidationRow] = []
    for interface_name, interface_result in interfaces.items():
        row = _format_single_interface_row(
            device_name, device_rack_pos, interface_name, interface_result
        )
        if row is None:
            continue
        if _should_skip_link_down_dedup(row, dedup):
            continue
        if _should_skip_no_neighbor(row, ignore_no_neighbor):
            continue
        results.append(row)
    return results


def _generate_notes(results: list[CableValidationRow]) -> list[str]:
    """Generate notes based on patterns in validation results.

    Args:
        results: List of CableValidationRow

    Returns:
        List of note strings
    """
    notes = []
    for row in results:
        issue = row.issue or ""
        if "*" in issue:
            notes.append(
                "*Either the expected MAC in our database is wrong"
                ", or this link is not cabled correctly."
            )
            break
    return notes


def _format_results_markdown(
    results: list[CableValidationRow],
    csv_exclude_columns: Iterable[str],
    markdown_exclude_columns: Iterable[str],
    empty_message: str = "No invalid cabling found.",
    max_display_results: int = 1000,
    export: str = "csv",
    summary_devices: dict[str, CableValidationResultData] | None = None,
) -> str:
    """Format validation results into markdown with an export link, table, and notes.

    Args:
        results: List of CableValidationRow
        csv_exclude_columns: Columns to exclude from CSV export (unused; model uses CSV_EXCLUDE_COLUMNS)
        markdown_exclude_columns: Columns to exclude from markdown table (unused; model uses MARKDOWN_EXCLUDE_COLUMNS)
        empty_message: Message to display when there are no results
        max_display_results: Maximum number of results to display in table
        export: "csv" for a single CSV link, "xlsx" for a two-tab Excel download
        summary_devices: All queried switches, used to seed the Excel summary tab so
            healthy switches appear with zero counts

    Returns:
        Formatted markdown string
    """
    if not results:
        return empty_message

    markdown_rows = [r.to_markdown() for r in results]
    export_link = (
        _generate_xlsx_link(results, summary_devices)
        if export == "xlsx"
        else _generate_csv_link(results)
    )

    markdown = f"{export_link}\n"

    if len(results) > max_display_results:
        export_name = "Excel" if export == "xlsx" else "CSV"
        markdown += (
            f"Too many results to display ({len(results)} errors), "
            f"please export to {export_name} to view.\n"
        )
    else:
        markdown += str(
            markdown_table(markdown_rows).set_params(quote=False, row_sep="markdown").get_markdown()
        )

    notes = _generate_notes(results)
    if notes:
        markdown += "\n\n" + "\n".join(notes)

    return markdown


def _classify_issue(issue: str | None) -> str:
    """Bucket a row's issue into a host-summary category.

    Returns one of: "missing", "miscabled", "unexpected", "other". "missing"
    covers links that should exist but don't.
    """
    if issue in (LINK_DOWN_MSG, LINK_UP_NO_NEIGHBOR_MSG):
        return "missing"
    if issue and issue.startswith(INCORRECT_CABLING_PREFIX):
        return "miscabled"
    if issue == UNEXPECTED_CONNECTION_MSG:
        return "unexpected"
    return "other"


def _build_host_summary(
    results: list[CableValidationRow],
    devices: dict[str, CableValidationResultData] | None = None,
) -> list[dict[str, Any]]:
    """Aggregate per-switch issue counts for the summary tab."""
    summary: dict[str, dict[str, Any]] = {}

    def _get_bucket(host: str, rack: str | None) -> dict[str, Any]:
        bucket = summary.get(host)
        if bucket is None:
            bucket = {
                "Host": host,
                "Rack": rack,
                "Missing Cables": 0,
                "Miscabled": 0,
                "Unexpected Connections": 0,
                "Total Issues": 0,
            }
            summary[host] = bucket
        elif not bucket["Rack"] and rack:
            bucket["Rack"] = rack
        return bucket

    # Seed a row for every queried switch so healthy ones still appear (0 counts).
    # Match the host key to how rows are keyed (_format_device_result_row lowercases
    # the device name) so a switch is never listed twice.
    for name, data in (devices or {}).items():
        device = data.device
        host = device.name.lower() if device and device.name else name
        rack = _rack_position_str(device.rack, device.position) if device else None
        _get_bucket(host, rack)

    for row in results:
        bucket = _get_bucket(row.start_device or "(unknown)", row.start_rack)
        bucket["Total Issues"] += 1
        category = _classify_issue(row.issue)
        if category == "missing":
            bucket["Missing Cables"] += 1
        elif category == "miscabled":
            bucket["Miscabled"] += 1
        elif category == "unexpected":
            bucket["Unexpected Connections"] += 1

    return sorted(
        summary.values(),
        key=lambda r: (-r["Missing Cables"], -r["Total Issues"], r["Host"]),
    )


def _style_excel_sheet(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
    """Apply uniform column widths and an autofilter to a worksheet."""
    worksheet = writer.sheets[sheet_name]
    for col_idx in range(len(df.columns)):
        worksheet.column_dimensions[get_column_letter(col_idx + 1)].width = 22
    if len(df) > 0:
        last_col = get_column_letter(len(df.columns))
        worksheet.auto_filter.ref = f"A1:{last_col}{len(df) + 1}"


def _build_cable_validation_workbook(
    results: list[CableValidationRow],
    devices: dict[str, CableValidationResultData] | None = None,
) -> bytes:
    """Build a two-tab .xlsx: full cable issue detail plus a per-switch summary."""
    detail_rows = [r.to_csv_dict() for r in results]
    detail_columns = (
        list(detail_rows[0].keys())
        if detail_rows
        else list(CableValidationRow().to_csv_dict().keys())  # type: ignore[call-arg]
    )
    detail_df = pd.DataFrame(detail_rows, columns=detail_columns)
    summary_df = pd.DataFrame(
        _build_host_summary(results, devices), columns=list(HOST_SUMMARY_COLUMNS)
    )

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        detail_df.to_excel(writer, sheet_name=DETAIL_SHEET_NAME, index=False)
        _style_excel_sheet(writer, DETAIL_SHEET_NAME, detail_df)
        summary_df.to_excel(writer, sheet_name=HOST_SUMMARY_SHEET_NAME, index=False)
        _style_excel_sheet(writer, HOST_SUMMARY_SHEET_NAME, summary_df)
    return buffer.getvalue()


def _generate_xlsx_link(
    results: list[CableValidationRow],
    devices: dict[str, CableValidationResultData] | None = None,
) -> str:
    """Generate an Excel download link with detail and host-summary tabs."""
    workbook = _build_cable_validation_workbook(results, devices)
    b64_xlsx = base64.b64encode(workbook).decode("utf-8")
    return f"[Download Excel](data:{EXCEL_MIME_TYPE};base64,{b64_xlsx})"


def _generate_csv_link(results: list[CableValidationRow]) -> str:
    """Generate a CSV export link from validation results.

    Args:
        results: List of CableValidationRow

    Returns:
        CSV data URI string for download (markdown link)
    """
    if not results:
        return "[Export to CSV](data:text/csv;base64,)"

    csv_rows = [r.to_csv_dict() for r in results]
    result_csv = io.StringIO()
    writer = csv.DictWriter(result_csv, csv_rows[0].keys())
    writer.writeheader()
    writer.writerows(csv_rows)
    csv_string = result_csv.getvalue()
    b64_csv = base64.b64encode(csv_string.encode("utf-8"))
    return f"[Export to CSV](data:text/csv;base64,{b64_csv.decode()})"


__all__ = [
    "_as_mac_address",
    "_build_actual_neighbor_data",
    "_build_cable_validation_workbook",
    "_build_host_summary",
    "_check_device_mac_match",
    "_check_interface_mac_match",
    "_check_link_state",
    "_check_lldp_validity",
    "_check_mac_validity",
    "_check_standard_lldp_match",
    "_classify_issue",
    "_display_or_macs",
    "_escape_formula",
    "_find_unexpected_neighbors",
    "_format_device_result_row",
    "_format_results_markdown",
    "_format_single_interface_row",
    "_generate_csv_link",
    "_generate_notes",
    "_generate_xlsx_link",
    "_mac_matches_with_offset",
    "_process_intended_interface",
    "_rack_position_str",
    "_row_link_down",
    "_row_link_up_no_neighbor",
    "_row_lldp_mismatch",
    "_row_mac_mismatch",
    "_row_unexpected_connection",
    "_should_skip_link_down_dedup",
    "_should_skip_no_neighbor",
    "_style_excel_sheet",
    "_validate_neighbor_has_required_fields",
    "format_mac",
    "is_mac_address",
]
