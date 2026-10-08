# SPDX-FileCopyrightText: Copyright (c) 2024-2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Excel-generation helpers for hardware-validation activities."""

import base64
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.utils import get_column_letter

from nv_config_manager_workflows.activities.hardware_validation.models import (
    CreateConsolidatedExcelInput,
    CreateConsolidatedExcelOutput,
    CreateExcelInput,
    CreateExcelOutput,
)


def create_excel(activity_input: CreateExcelInput) -> CreateExcelOutput:
    """Create an Excel export for one hardware-validation stage."""
    command_name = activity_input.command_name
    devices_data_and_results = activity_input.devices_data_and_results
    rows = []

    for result in devices_data_and_results.values():
        device_data = result["device_data"]
        command_result = result.get("command_result", {})

        if result.get("error"):
            continue

        if isinstance(device_data, dict):
            device_name = device_data.get("name", "N/A")
            rack_name = device_data.get("rack", "N/A") or "N/A"
            rack_position = device_data.get("position", "N/A") or "N/A"
        else:
            device_name = device_data.name
            rack_name = device_data.rack or "N/A"
            rack_position = device_data.position or "N/A"

        if command_name == "platform":
            for value in command_result.values():
                if isinstance(value, dict):
                    rows.append(
                        {
                            "device_name": device_name,
                            "rack_name": rack_name,
                            "rack_position": str(rack_position),
                            **_flatten_dict(value),
                        }
                    )
        else:
            for item_name, item_data in command_result.items():
                if isinstance(item_data, dict):
                    rows.append(
                        {
                            "device_name": device_name,
                            "rack_name": rack_name,
                            "rack_position": str(rack_position),
                            "item_name": item_name,
                            **_flatten_dict(item_data),
                        }
                    )

    if rows:
        df = _prepare_dataframe(rows)
        excel_buffer = BytesIO()
        with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
            sheet_name = command_name.capitalize()
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            _set_column_widths(writer.sheets[sheet_name], len(df.columns))
        row_count = len(df)
    else:
        df = _empty_dataframe()
        excel_buffer = BytesIO()
        with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
            df.to_excel(writer, index=False)
            _set_column_widths(writer.sheets["Sheet1"], len(df.columns))
        row_count = 0

    return CreateExcelOutput(
        excel_data=base64.b64encode(excel_buffer.getvalue()).decode("utf-8"),
        row_count=row_count,
    )


def create_consolidated_excel(
    activity_input: CreateConsolidatedExcelInput,
) -> CreateConsolidatedExcelOutput:
    """Create an Excel export with one worksheet per hardware-validation stage."""
    excel_buffer = BytesIO()
    worksheet_counts = {}
    total_row_count = 0

    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        for stage_name, devices_data_and_results in activity_input.stage_data.items():
            rows = []

            for result in devices_data_and_results.values():
                device_data = result["device_data"]
                command_result = result.get("command_result", {})

                if result.get("error"):
                    continue

                if isinstance(device_data, dict):
                    device_name = device_data.get("name", "N/A")
                    rack_name = device_data.get("rack", "N/A") or "N/A"
                    rack_position = device_data.get("position", "N/A") or "N/A"
                else:
                    device_name = device_data.name
                    rack_name = device_data.rack or "N/A"
                    rack_position = device_data.position or "N/A"

                if stage_name == "platform":
                    if isinstance(command_result, dict):
                        rows.append(
                            {
                                "device_name": device_name,
                                "rack_name": rack_name,
                                "rack_position": str(rack_position),
                                **_flatten_dict(command_result),
                            }
                        )
                elif isinstance(command_result, dict):
                    for item_name, item_data in command_result.items():
                        if isinstance(item_data, dict):
                            rows.append(
                                {
                                    "device_name": device_name,
                                    "rack_name": rack_name,
                                    "rack_position": str(rack_position),
                                    "item_name": item_name,
                                    **_flatten_dict(item_data),
                                }
                            )

            sheet_name = stage_name.capitalize()
            if rows:
                df = _prepare_dataframe(rows, include_item_name=True, sort=True)
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                worksheet = writer.sheets[sheet_name]
                _set_column_widths(worksheet, len(df.columns))
                worksheet.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
                worksheet_counts[sheet_name] = len(df)
                total_row_count += len(df)
            else:
                df = _empty_dataframe()
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                worksheet = writer.sheets[sheet_name]
                _set_column_widths(worksheet, len(df.columns))
                worksheet.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}1"
                worksheet_counts[sheet_name] = 0

    return CreateConsolidatedExcelOutput(
        excel_data=base64.b64encode(excel_buffer.getvalue()).decode("utf-8"),
        total_row_count=total_row_count,
        worksheet_counts=worksheet_counts,
    )


def _prepare_dataframe(
    rows: list[dict[str, Any]],
    *,
    include_item_name: bool = False,
    sort: bool = False,
) -> pd.DataFrame:
    """Normalize, order, and format a non-empty export dataframe."""
    df = pd.DataFrame(rows)
    for column in df.columns:
        if df[column].dtype not in ["int64", "int32", "float64", "float32"]:
            df[column] = df[column].fillna("N/A")

    priority_columns = ["device_name", "rack_name", "rack_position"]
    if include_item_name and "item_name" in df.columns:
        priority_columns.append("item_name")
    existing_priority = [column for column in priority_columns if column in df.columns]
    other_columns = sorted(column for column in df.columns if column not in priority_columns)
    df = df[existing_priority + other_columns]

    if sort:
        sort_columns = [column for column in ("rack_name", "rack_position") if column in df.columns]
        if sort_columns:
            df = df.sort_values(sort_columns)

    return df.rename(columns={column: _format_column_header(column) for column in df.columns})


def _empty_dataframe() -> pd.DataFrame:
    """Build the standard empty hardware-validation export dataframe."""
    df = pd.DataFrame(columns=["device_name", "rack_name", "rack_position", "No data available"])
    return df.rename(columns={column: _format_column_header(column) for column in df.columns})


def _set_column_widths(worksheet: Any, column_count: int) -> None:
    """Set the standard width on all worksheet columns."""
    for column_index in range(column_count):
        column_letter = get_column_letter(column_index + 1)
        worksheet.column_dimensions[column_letter].width = 20


def _format_column_header(header: str) -> str:
    """Format column header for better readability."""
    formatted = header.replace("_", " ").replace("-", " ")
    return " ".join(word.capitalize() for word in formatted.split())


def _flatten_dict(data: dict[str, Any], parent_key: str = "", sep: str = "_") -> dict[str, Any]:
    """Flatten a nested dictionary."""
    items: list[tuple[str, Any]] = []
    for key, value in data.items():
        new_key = f"{parent_key}{sep}{key}" if parent_key else key
        if isinstance(value, dict):
            items.extend(_flatten_dict(value, new_key, sep=sep).items())
        else:
            items.append((new_key, value))
    return dict(items)


__all__ = [
    "create_consolidated_excel",
    "create_excel",
]
