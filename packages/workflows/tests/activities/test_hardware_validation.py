# SPDX-FileCopyrightText: Copyright (c) 2024-2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Package-owned hardware validation activity contracts."""

import base64
import io
from collections.abc import Callable
from typing import cast
from unittest.mock import MagicMock

import pytest
from nv_config_manager_dcim.workflow_models import NetworkDeviceData
from openpyxl import load_workbook

from nv_config_manager_workflows.activities import hardware_validation
from nv_config_manager_workflows.activities.hardware_validation import (
    HARDWARE_VALIDATION_ACTIVITIES,
    CreateExcelInput,
    HardwareValidationInput,
    HardwareValidationOutput,
    create_excel_export,
    get_platform,
    get_platform_environment_fan,
    get_platform_environment_led,
    get_platform_environment_psu,
    get_platform_environment_voltage,
    get_platform_inventory,
    helpers,
)
from nv_config_manager_workflows.clients.device.base import NetworkConnection
from nv_config_manager_workflows.registration import activity_name
from nv_config_manager_workflows.runtime import configure_device_connection


def _device() -> NetworkDeviceData:
    return NetworkDeviceData.model_construct(
        id="device-1",
        name="leaf-1",
        rack="rack-1",
        position=42,
        role="leaf",
        site="site-1",
        device_type="switch",
        platform="cumulus-linux",
        primary_ip4="192.0.2.1",
        primary_ip6=None,
    )


def test_hardware_catalog_has_eight_unique_activities() -> None:
    assert len(HARDWARE_VALIDATION_ACTIVITIES) == 8
    assert len({activity_name(item) for item in HARDWARE_VALIDATION_ACTIVITIES}) == 8


def test_private_excel_helpers_are_not_exported() -> None:
    assert "_flatten_dict" not in helpers.__all__
    assert "_format_column_header" not in helpers.__all__
    assert not hasattr(hardware_validation, "_flatten_dict")
    assert not hasattr(hardware_validation, "_format_column_header")


def test_excel_helpers_remain_directly_testable() -> None:
    assert helpers._flatten_dict({"system": {"model_name": "SN5600"}}) == {
        "system_model_name": "SN5600"
    }
    assert helpers._format_column_header("system_model-name") == "System Model Name"


def test_platform_query_uses_runtime_device_connection() -> None:
    connection = MagicMock()
    connection.get_platform.return_value = {"system": {"model": "SN5600"}}
    configure_device_connection(lambda _device: cast(NetworkConnection, connection))

    result = get_platform(HardwareValidationInput(device_data=_device()))

    assert result.info == {"system": {"model": "SN5600"}}
    connection.get_platform.assert_called_once_with()


@pytest.mark.parametrize(
    ("activity_function", "connection_method"),
    [
        (get_platform_environment_fan, "get_platform_environment_fan"),
        (get_platform_environment_led, "get_platform_environment_led"),
        (get_platform_environment_psu, "get_platform_environment_psu"),
        (get_platform_environment_voltage, "get_platform_environment_voltage"),
        (get_platform_inventory, "get_platform_inventory"),
    ],
)
def test_hardware_query_uses_matching_connection_method(
    activity_function: Callable[[HardwareValidationInput], HardwareValidationOutput],
    connection_method: str,
) -> None:
    connection = MagicMock()
    getattr(connection, connection_method).return_value = {"component": {"state": "ok"}}
    configure_device_connection(lambda _device: cast(NetworkConnection, connection))

    result = activity_function(HardwareValidationInput(device_data=_device()))

    assert result.info == {"component": {"state": "ok"}}
    getattr(connection, connection_method).assert_called_once_with()


def test_single_command_export_is_a_real_workbook() -> None:
    result = create_excel_export(
        CreateExcelInput(
            command_name="fan",
            devices_data_and_results={
                "leaf-1": {
                    "device_data": _device(),
                    "command_result": {"fan-1": {"state": "ok"}},
                }
            },
        )
    )

    workbook = load_workbook(io.BytesIO(base64.b64decode(result.excel_data)))
    assert result.row_count == 1
    assert workbook.sheetnames == ["Fan"]
    assert [cell.value for cell in workbook["Fan"][1]][:4] == [
        "Device Name",
        "Rack Name",
        "Rack Position",
        "Item Name",
    ]


def test_hardware_single_workbook_contract_is_frozen() -> None:
    """Single-stage exports retain flattening, ordering, null, and workbook formatting."""
    result = hardware_validation.create_excel_export(
        hardware_validation.CreateExcelInput(
            command_name="fan",
            devices_data_and_results={
                "device-b": {
                    "device_data": {
                        "name": "leaf-b",
                        "rack": "rack-b",
                        "position": 2,
                    },
                    "command_result": {
                        "FAN2": {
                            "state": None,
                            "metrics": {"speed": 2000},
                        }
                    },
                },
                "failed": {
                    "device_data": {"name": "leaf-failed"},
                    "error": "unreachable",
                },
                "device-a": {
                    "device_data": {
                        "name": "leaf-a",
                        "rack": None,
                        "position": None,
                    },
                    "command_result": {
                        "FAN1": {
                            "state": "ok",
                            "metrics": {"speed": 1000},
                            "temperature": 42.5,
                        }
                    },
                },
            },
        )
    )

    assert result.row_count == 2
    assert result.excel_data.startswith("UEsDB")
    workbook = load_workbook(io.BytesIO(base64.b64decode(result.excel_data)))
    assert workbook.sheetnames == ["Fan"]
    worksheet = workbook["Fan"]
    assert list(worksheet.values) == [
        (
            "Device Name",
            "Rack Name",
            "Rack Position",
            "Item Name",
            "Metrics Speed",
            "State",
            "Temperature",
        ),
        ("leaf-b", "rack-b", "2", "FAN2", 2000, "N/A", None),
        ("leaf-a", "N/A", "N/A", "FAN1", 1000, "ok", 42.5),
    ]
    assert [worksheet.column_dimensions[column].width for column in "ABCDEFG"] == [
        20.0,
        20.0,
        20.0,
        20.0,
        20.0,
        20.0,
        20.0,
    ]
    assert worksheet.auto_filter.ref is None


def test_hardware_consolidated_workbook_contract_is_frozen() -> None:
    """Consolidated exports retain sheet/count/order/header/filter/base64 contracts."""
    result = hardware_validation.create_consolidated_excel_export(
        hardware_validation.CreateConsolidatedExcelInput(
            stage_data={
                "platform": {
                    "device-b": {
                        "device_data": {
                            "name": "leaf-b",
                            "rack": "rack-b",
                            "position": 2,
                        },
                        "command_result": {
                            "product-name": "SN5600",
                            "system": {"memory-gb": 32},
                        },
                    }
                },
                "fan": {
                    "device-b": {
                        "device_data": {
                            "name": "leaf-b",
                            "rack": "rack-b",
                            "position": 2,
                        },
                        "command_result": {"FAN2": {"state": None, "speed": 2000}},
                    },
                    "device-a": {
                        "device_data": {
                            "name": "leaf-a",
                            "rack": "rack-a",
                            "position": 10,
                        },
                        "command_result": {"FAN1": {"state": "ok", "speed": 1000}},
                    },
                },
                "empty": {
                    "failed": {
                        "device_data": {"name": "leaf-failed"},
                        "error": "unreachable",
                    }
                },
            }
        )
    )

    assert result.total_row_count == 3
    assert result.worksheet_counts == {"Platform": 1, "Fan": 2, "Empty": 0}
    assert result.excel_data.startswith("UEsDB")
    workbook = load_workbook(io.BytesIO(base64.b64decode(result.excel_data)))
    assert workbook.sheetnames == ["Platform", "Fan", "Empty"]
    assert list(workbook["Platform"].values) == [
        (
            "Device Name",
            "Rack Name",
            "Rack Position",
            "Product Name",
            "System Memory Gb",
        ),
        ("leaf-b", "rack-b", "2", "SN5600", 32),
    ]
    assert list(workbook["Fan"].values) == [
        ("Device Name", "Rack Name", "Rack Position", "Item Name", "Speed", "State"),
        ("leaf-a", "rack-a", "10", "FAN1", 1000, "ok"),
        ("leaf-b", "rack-b", "2", "FAN2", 2000, "N/A"),
    ]
    assert list(workbook["Empty"].values) == [
        ("Device Name", "Rack Name", "Rack Position", "No Data Available")
    ]
    assert workbook["Platform"].auto_filter.ref == "A1:E2"
    assert workbook["Fan"].auto_filter.ref == "A1:F3"
    assert workbook["Empty"].auto_filter.ref == "A1:D1"
    for worksheet in workbook.worksheets:
        assert all(
            worksheet.column_dimensions[column].width == 20.0
            for column in tuple(worksheet.column_dimensions)
        )
